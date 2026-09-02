from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


BLUE = RGBColor(31, 78, 121)
GREEN = RGBColor(31, 122, 79)
RED = RGBColor(192, 0, 0)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, size=10.5, *, bold=False, color=BLACK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_text(cell, text, *, size=9.0, bold=False, color=BLACK, alignment=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_response(cell, paragraphs):
    for paragraph in list(cell.paragraphs):
        if not paragraph.text.strip():
            paragraph._element.getparent().remove(paragraph._element)
    label = cell.add_paragraph()
    label.paragraph_format.space_before = Pt(8)
    label.paragraph_format.space_after = Pt(3)
    run = label.add_run("Respuesta:")
    set_run_font(run, size=10.5, bold=True, color=BLUE)
    for text in paragraphs:
        paragraph = cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def replace_in_runs(paragraph, replacements):
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def fill_11(template: Path, output: Path):
    shutil.copy2(template, output)
    doc = Document(output)

    details = doc.tables[1]
    set_cell_text(
        details.cell(0, 0),
        "Escuela de Informática y Telecomunicaciones",
        size=10.5,
        bold=True,
        color=BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_text(details.cell(1, 1), "Benjamín Olmedo", size=10.5)
    set_cell_text(details.cell(2, 1), "Ingeniería en Informática", size=10.5)
    set_cell_text(details.cell(3, 1), "2023", size=10.5)

    competencies = [
        (
            "Analizar procesos y proponer soluciones informáticas",
            "AD",
            "Puedo comprender necesidades y proponer soluciones creativas; debo reforzar la validación formal de requisitos con usuarios.",
            GREEN,
        ),
        (
            "Desarrollar soluciones de software",
            "AD",
            "Es una de mis principales fortalezas. Me siento seguro construyendo interfaces y funcionalidades, especialmente para entornos web y móviles.",
            GREEN,
        ),
        (
            "Construir programas usando buenas prácticas",
            "AD",
            "Puedo transformar requerimientos en código funcional y organizado; todavía puedo profundizar en patrones y mantenibilidad a gran escala.",
            GREEN,
        ),
        (
            "Realizar pruebas de productos y procesos",
            "DA",
            "Realizo verificaciones funcionales, pero necesito fortalecer automatización, cobertura y documentación sistemática de pruebas.",
            RED,
        ),
        (
            "Construir modelos arquitectónicos de soluciones",
            "DA",
            "Comprendo la separación de componentes y responsabilidades, aunque necesito más experiencia diseñando arquitecturas escalables y seguras.",
            RED,
        ),
        (
            "Gestionar proyectos informáticos",
            "AD",
            "Me interesa especialmente planificar, coordinar personas y tomar decisiones. He aplicado Scrum y trabajo colaborativo en proyectos académicos.",
            GREEN,
        ),
        (
            "Comunicarse en inglés en contextos laborales",
            "AD",
            "Puedo comprender documentación y comunicar ideas técnicas en inglés; continuaré ampliando vocabulario profesional y fluidez oral.",
            GREEN,
        ),
        (
            "Elaborar proyectos innovadores que agreguen valor",
            "AD",
            "La creatividad y la visión de producto son fortalezas que quiero convertir en emprendimientos tecnológicos sostenibles y útiles.",
            GREEN,
        ),
    ]
    rating_col = {"ED": 1, "AD": 2, "DA": 3, "DP": 4, "DNL": 5}
    grid = doc.tables[2]
    for row_index, (name, rating, comment, color) in enumerate(competencies, start=2):
        set_cell_text(grid.cell(row_index, 0), name, size=8.2, bold=True, color=color)
        for col in range(1, 6):
            set_cell_text(grid.cell(row_index, col), "X" if col == rating_col[rating] else "", size=10.0, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(grid.cell(row_index, 6), comment, size=8.0)

    doc.core_properties.title = "Autoevaluación de Competencias Fase 1 - Benjamín Olmedo"
    doc.core_properties.subject = "APT122 - Evidencia individual 1.1"
    doc.core_properties.author = "Benjamín Olmedo"
    doc.save(output)


def fill_12(template: Path, output: Path):
    shutil.copy2(template, output)
    doc = Document(output)

    add_response(
        doc.tables[2].cell(1, 0),
        [
            "Las áreas que más me han gustado son Desarrollo Web, Desarrollo Móvil y Gestión de Proyectos. En desarrollo web disfruté convertir una idea en una interfaz funcional y accesible desde distintos dispositivos. En desarrollo móvil me motivó diseñar experiencias pensadas para el uso cotidiano. En gestión de proyectos me interesó la planificación, la organización del equipo, la priorización y la toma de decisiones.",
            "También valoro la experiencia inicial de Cisco vinculada con Python y un ramo electivo de proyectos en el que trabajamos con Scrum. Aunque no recuerdo el nombre oficial de todas las certificaciones obtenidas, considero que tienen valor porque respaldan conocimientos concretos, entregan una referencia externa de aprendizaje y complementan la experiencia práctica desarrollada durante la carrera.",
        ],
    )

    add_response(
        doc.tables[3].cell(1, 0),
        [
            "Mis fortalezas principales son el desarrollo de software, la gestión de proyectos, la creatividad, el trabajo en equipo y el inglés. Me siento especialmente seguro participando en la definición de una solución, construyendo su frontend, proponiendo mejoras y coordinando tareas con otras personas.",
            "Las competencias que necesito fortalecer son la arquitectura de soluciones escalables, las pruebas automatizadas, la ciberseguridad y la formalización de requisitos. Puedo aplicar sus conceptos centrales, pero requiero más práctica para utilizarlos con profundidad en productos reales que deban crecer y mantenerse en el tiempo.",
        ],
    )

    add_response(
        doc.tables[4].cell(1, 0),
        [
            "Mis principales intereses profesionales son el desarrollo de productos digitales, el emprendimiento tecnológico y la dirección de proyectos. Quiero comenzar creando software y luego asumir responsabilidades de estrategia, gestión y liderazgo.",
            "Las competencias más relacionadas son desarrollo de software, análisis de requerimientos, innovación, gestión de proyectos, trabajo en equipo e inglés técnico. Necesito fortalecer arquitectura, seguridad y calidad para convertir prototipos en productos escalables.",
            "En cinco años me gustaría dirigir proyectos propios exitosos: comenzar construyendo el software y luego contratar un equipo que me permita dedicarme a la dirección y al crecimiento del producto como fundador.",
        ],
    )

    add_response(
        doc.tables[5].cell(1, 0),
        [
            "La propuesta vigente se llama Fix & Go. Es una plataforma web y móvil para conectar a personas que necesitan resolver servicios técnicos del hogar con profesionales cercanos y verificables. La entrada del proyecto se concentra en gas, electricidad y agua; otras categorías quedan como expansión posterior a la validación.",
            "El producto completo considera perfiles diferenciados, validación de credenciales, búsqueda, solicitudes, cotizaciones, mensajería interna, pagos seguros, seguimiento y reputación. Para APT122 construiremos un MVP con datos ficticios e integraciones sensibles simuladas, de forma que podamos probar el recorrido sin procesar dinero, documentos o direcciones reales.",
            "Fix & Go aporta a mi desarrollo profesional porque combina frontend móvil y web, experiencia de usuario, análisis, arquitectura, seguridad, pruebas y gestión. Mi aporte inicial se enfocará en frontend, prototipado y consistencia de la experiencia. La distribución definitiva de backend, datos y calidad se acordará con Daniel Baeza y Felipe Arce según las fortalezas de cada integrante.",
        ],
    )

    doc.core_properties.title = "Diario de Reflexión Fase 1 - Benjamín Olmedo"
    doc.core_properties.subject = "APT122 - Evidencia individual 1.2"
    doc.core_properties.author = "Benjamín Olmedo"
    doc.save(output)


def fill_13(path: Path):
    doc = Document(path)
    replacements = {
        4: "FIX & GO",
        9: (
            "Fix & Go es una plataforma web y móvil que conecta a personas que necesitan servicios técnicos del hogar con profesionales cercanos y verificables. El proyecto responde a una contratación fragmentada entre recomendaciones, redes sociales y contactos con poca información comparable sobre especialidad, credenciales, cobertura, disponibilidad, precio o reputación. La entrada se concentra en gas, electricidad y agua. El MVP académico permitirá buscar, revisar perfiles, crear una solicitud, comparar cotizaciones, seguir un servicio y registrar una valoración. Verificación documental, mensajería, pagos y geolocalización se representarán mediante datos ficticios y simulaciones explícitas. El trabajo será incremental, con requisitos versionados, prototipos, control de cambios, pruebas y validación con usuarios potenciales. La evidencia permitirá evaluar si el recorrido mejora la claridad, confianza y trazabilidad del proceso actual."
        ),
        10: "Palabras clave: servicios técnicos, profesionales verificados, aplicación móvil, plataforma web, marketplace, experiencia de usuario.",
        12: (
            "Fix & Go is a web and mobile platform that connects people who need home technical services with nearby, verifiable professionals. The project addresses a fragmented hiring process based on referrals, social media, and contacts with little comparable information about expertise, credentials, coverage, availability, price, or reputation. The initial categories are gas, electricity, and water. The academic MVP will cover discovery, profiles, service requests, quote comparison, job tracking, and reviews. Credential verification, messaging, payments, and production geolocation will be represented through fictional data and clearly labelled simulations. The team will work incrementally through versioned requirements, prototypes, source control, testing, and validation with potential users. The evidence will assess whether the flow improves clarity, trust, and traceability."
        ),
        13: "Keywords: technical services, verified professionals, mobile application, web platform, marketplace, user experience.",
        15: (
            "Fix & Go propone centralizar la búsqueda y contratación de profesionales para servicios de gas, electricidad y agua. Actualmente, muchas personas dependen de recomendaciones informales o publicaciones dispersas, lo que dificulta comparar perfiles, verificar credenciales, conocer cobertura y disponibilidad, entender el precio y mantener trazabilidad sobre el servicio."
        ),
        16: (
            "La solución considera tres actores principales: el cliente busca, solicita, contrata y evalúa; el profesional administra perfil, credenciales, cobertura, cotizaciones y trabajos; y el administrador revisa verificaciones, incidentes y auditoría. El flujo parte con la necesidad y termina con el cierre, pago representado y evaluación verificable."
        ),
        17: (
            "La propuesta es relevante para la informática porque integra análisis de requerimientos, experiencia móvil, arquitectura, modelado de datos, seguridad y privacidad, desarrollo de interfaces, integración, control de versiones y pruebas. Además, permite aplicar gestión de proyectos e innovación a un producto digital con posibilidad de evolución comercial (Duoc UC, s.f.)."
        ),
        23: (
            "Mis intereses profesionales se concentran en el desarrollo de software, la creación de productos digitales y la dirección de proyectos tecnológicos. Me interesa iniciar participando directamente en la construcción de una solución y, con el crecimiento del producto, asumir la estrategia, la coordinación y el liderazgo de un equipo. Fix & Go conecta esos intereses porque exige transformar una necesidad cotidiana en una experiencia móvil clara, verificable y con potencial de negocio."
        ),
        24: (
            "Mi aporte individual: Lideraré principalmente el frontend móvil, el prototipado y la consistencia de la experiencia de usuario. También aportaré creatividad, organización del trabajo, comunicación en inglés técnico y coordinación de decisiones. La asignación definitiva de backend, datos y calidad se acordará con Daniel Baeza y Felipe Arce según las fortalezas de cada integrante."
        ),
        26: (
            "Fix & Go es factible si el equipo limita la primera entrega al recorrido esencial y simula las integraciones de mayor riesgo. Los tres integrantes cuentan con computadores, GitHub y experiencia académica en desarrollo y colaboración. El MVP usará datos ficticios y no procesará pagos, credenciales ni direcciones reales. La construcción incremental permitirá validar primero el problema, luego la experiencia web/móvil y finalmente la coherencia del flujo completo."
        ),
        29: "Objetivo general. Diseñar, construir y validar durante APT122 un MVP web y móvil de Fix & Go que represente la búsqueda, contratación y seguimiento de profesionales verificados para gas, electricidad y agua mediante una experiencia simple, segura en su simulación y trazable.",
        30: "Objetivo específico 1. Levantar y formalizar las necesidades de clientes y profesionales, definiendo actores, reglas, estados, riesgos y criterios de aceptación para el flujo principal.",
        31: "Objetivo específico 2. Modelar arquitectura y datos escalables para usuarios, credenciales, categorías, cobertura, solicitudes, cotizaciones, servicios, pagos simulados y valoraciones.",
        32: "Objetivo específico 3. Implementar una experiencia web y móvil accesible para buscar, comparar perfiles, solicitar, cotizar, seguir el servicio y evaluar.",
        33: "Objetivo específico 4. Evaluar el MVP mediante pruebas funcionales, revisión de interfaz y sesiones con usuarios potenciales, utilizando los hallazgos para priorizar mejoras.",
        35: "Se utilizará una metodología incremental y colaborativa, apoyada en prácticas de Scrum adaptadas al contexto académico. Cada ciclo comenzará con un objetivo y criterios de aceptación, continuará con diseño e implementación, y terminará con verificación, evidencia y ajuste (Schwaber & Sutherland, 2020).",
        36: "Descubrir y validar necesidades mediante entrevistas breves a personas que contratan servicios y a profesionales independientes.",
        37: "Especificar historias de usuario, reglas, estados, riesgos y criterios de aceptación en documentación versionada.",
        38: "Diseñar el modelo de datos, la arquitectura móvil, la navegación y el prototipo antes de integrar funcionalidades de mayor riesgo.",
        39: "Construir incrementos pequeños en ramas separadas, revisados mediante pull request antes de integrarlos a main.",
        40: "Verificar componentes y flujos críticos; revisar accesibilidad, adaptación a pantallas, consistencia de datos, privacidad y manejo de errores.",
        41: "Validar con usuarios potenciales, registrar éxito, tiempos, dudas y errores, y priorizar mejoras por impacto y frecuencia.",
        44: "Recursos transversales: computadores personales, GitHub, entornos de desarrollo web y móvil, documentación, datos de demostración e instrumentos de prueba. La elección final del stack y la distribución de backend, datos y calidad se acordarán entre los tres integrantes antes de implementar.",
        48: "La definición incorpora los indicadores 1.1, 1.3, 2.1, 2.2, 3.1, 3.2, 4.1, 4.2 y 4.3. En esta fase se evidencia principalmente el diseño de pruebas, la planificación, el modelo escalable y la arquitectura de integración. La implementación, el control y la validación se demostrarán progresivamente mediante GitHub, el MVP y los reportes del equipo.",
        49: "Pruebas: estrategia para búsqueda, perfiles, creación de solicitudes, respuesta, estados y manejo de errores.",
        50: "Gestión: roadmap incremental, backlog, definición de terminado, GitHub y seguimiento de riesgos.",
        51: "Datos: usuarios, profesionales, credenciales, categorías, cobertura, solicitudes, cotizaciones, servicios y valoraciones con relaciones coherentes.",
        52: "Desarrollo: frontend web/móvil, lógica de aplicación, dominio e infraestructura integrados mediante interfaces y adaptadores definidos.",
        57: (
            "Fix & Go is a relevant and feasible Capstone project because it addresses a common technical-services problem through a bounded web and mobile MVP. The proposal connects trust, project management, data modeling, software development, privacy, security, and validation. Its strongest point is the traceable flow from discovering a verifiable professional to closing and reviewing a service. The next stage must validate assumptions with clients and professionals, agree on responsibilities, select the stack, and transform the specification into a tested prototype."
        ),
        59: (
            "This self-assessment connected my strengths in software development, creativity, teamwork, project management, and English with my goal of becoming a technology founder. Fix & Go lets me begin with frontend development while learning to manage scope and coordinate a growing product. I must strengthen architecture, automated testing, and security. In the next phase, I will lead the mobile interface, document decisions, and help define responsibilities with Daniel and Felipe."
        ),
        62: "Duoc UC. (s.f.). Descripción del perfil de egreso: Ingeniería en Informática [Documento académico].",
        63: "Fix & Go. (2026). Fuente oficial consolidada del Proyecto APT [Documento interno].",
        64: "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide. Scrum Guides.",
        65: "",
        66: "",
        67: "",
    }
    for index, text in replacements.items():
        replace_paragraph(doc.paragraphs[index], text)

    identification = [
        (0, 1, "Benjamín Olmedo"),
        (1, 1, "Benjamín Olmedo · Daniel Baeza · Felipe Arce · Equipo por definir"),
        (2, 1, "APT122 · Proyecto APT"),
        (3, 1, "005D"),
        (4, 1, "Fix & Go"),
        (5, 1, "2 de septiembre de 2026"),
    ]
    for row, col, text in identification:
        replace_paragraph(doc.tables[1].cell(row, col).paragraphs[0], text)

    replace_paragraph(
        doc.tables[2].cell(0, 0).paragraphs[-1],
        "El alcance evaluable es un MVP web y móvil del flujo buscar → revisar perfil → solicitar → cotizar → seguir servicio → evaluar. Pagos, mensajería, verificación y geolocalización se simulan con datos ficticios; la operación comercial real queda para una etapa posterior.",
    )

    competency_rows = [
        ("Pruebas y certificación", "1.1 y 1.3", "Diseñar pruebas del recorrido de búsqueda, cotización, servicio y evaluación, y convertir hallazgos en mejoras."),
        ("Gestión de proyectos informáticos", "2.1 y 2.2", "Planificar alcance, responsables, recursos, riesgos e hitos; controlar el avance con evidencia versionada."),
        ("Modelos de datos escalables", "3.1 y 3.2", "Diseñar entidades para usuarios, credenciales, solicitudes, cotizaciones, servicios, pagos y valoraciones."),
        ("Desarrollo de soluciones", "4.1, 4.2 y 4.3", "Construir e integrar experiencias web y móvil mediante técnicas sistemáticas de desarrollo y mantenimiento."),
    ]
    replace_paragraph(doc.tables[3].cell(0, 2).paragraphs[0], "Aplicación en Fix & Go")
    for row_index, values in enumerate(competency_rows, start=1):
        for col_index, value in enumerate(values):
            replace_paragraph(doc.tables[3].cell(row_index, col_index).paragraphs[0], value)

    feasibility = [
        ("Tiempo", "Plan semestral por fases; priorización del flujo central antes de módulos secundarios."),
        ("Materiales", "Tres computadores, GitHub, herramientas gratuitas de diseño y desarrollo móvil, y datos de demostración."),
        ("Factores externos", "Disponibilidad de clientes y profesionales independientes para entrevistas y validación."),
        ("Riesgo de alcance", "Simular pagos, mensajería, verificación y geolocalización; excluir operación comercial real."),
        ("Riesgo técnico y privacidad", "Prototipo temprano, permisos mínimos, datos ficticios, revisión por pares y pruebas incrementales."),
        ("Riesgo de coordinación", "Definir responsabilidades, mantener backlog visible, usar ramas y revisar cambios antes de integrarlos."),
    ]
    for row_index, values in enumerate(feasibility, start=1):
        for col_index, value in enumerate(values):
            replace_paragraph(doc.tables[4].cell(row_index, col_index).paragraphs[0], value)

    plan = [
        ("1. Definición", "Semanas 1-3", "Problema, usuarios, alcance y criterios", "Equipo completo", "Idea inicial y perfil de egreso", "Supuestos sin validar"),
        ("2. Producto y UX", "Semanas 4-5", "Flujos y prototipo web/móvil", "Benjamín lidera frontend; equipo valida", "Referencias y entrevistas", "Flujo demasiado amplio"),
        ("3. Arquitectura y datos", "Semanas 6-7", "Stack, modelo, repositorios y datos demo", "Distribución técnica por acordar", "GitHub y documentación", "Decisiones técnicas tardías"),
        ("4. Flujo vertical", "Semanas 8-10", "Búsqueda, perfil, solicitud y cotización", "Desarrollo compartido", "Criterios de aceptación", "Integración de estados"),
        ("5. Seguimiento", "Semanas 11-12", "Servicio, mensajes y valoración simulados", "Desarrollo compartido", "Flujo principal operativo", "Integraciones simuladas"),
        ("6. Calidad", "Semanas 13-14", "Pruebas, accesibilidad y correcciones", "Equipo completo", "Datos demo y checklist", "Tiempo de corrección limitado"),
        ("7. Validación", "Semanas 15-16", "Sesiones, resultados y presentación", "Equipo completo", "Acceso a usuarios potenciales", "Disponibilidad de participantes"),
    ]
    for row_index, values in enumerate(plan, start=1):
        for col_index, value in enumerate(values):
            replace_paragraph(doc.tables[5].cell(row_index, col_index).paragraphs[0], value)

    evidence = [
        ("Especificación del producto", "Problema, usuarios, requisitos, dominio, UX y decisiones", "Demuestra definición disciplinar y trazabilidad antes de construir."),
        ("Repositorio GitHub", "Commits, ramas, revisiones y documentación versionada", "Demuestra colaboración, control de cambios y evolución del producto."),
        ("Prototipo/MVP web y móvil", "Flujo de contratación ejecutable", "Evidencia integración de componentes y cumplimiento funcional."),
        ("Plan y tablero de trabajo", "Actividades, responsables, hitos, riesgos y estado", "Permite controlar avance y justificar decisiones de alcance."),
        ("Pruebas y reportes", "Casos automatizados, checklist manual e incidencias", "Demuestra verificación y mejoras basadas en resultados."),
        ("Registro de validación", "Tareas, tiempos, éxito, dudas y comentarios", "Contrasta la propuesta con clientes y profesionales potenciales."),
        ("Presentación final", "Problema, solución, demostración, resultados y próximos pasos", "Comunica de forma verificable el resultado y el aprendizaje."),
    ]
    for row_index, values in enumerate(evidence, start=1):
        for col_index, value in enumerate(values):
            replace_paragraph(doc.tables[6].cell(row_index, col_index).paragraphs[0], value)

    rubric_evidence = [
        "Se explica el problema, la solución móvil, los usuarios iniciales y su relación con el desarrollo de software.",
        "Se relacionan pruebas, gestión, datos y desarrollo con actividades concretas de Fix & Go.",
        "Se explicitan intereses de emprendimiento, fortalezas y aporte individual en frontend.",
        "Se delimitan tiempo, recursos, dependencias, riesgos técnicos, privacidad y medidas de respuesta.",
        "Se formula un objetivo general y cuatro objetivos específicos claros y verificables.",
        "Se propone trabajo incremental, control de versiones, pruebas y validación con usuarios.",
        "El plan contiene actividades, duración, responsables provisionales, facilitadores y obstáculos.",
        "Cada evidencia se vincula con actividades y resultados esperados del proyecto.",
        "El informe utiliza referencias académicas y una redacción formal y coherente.",
        "Incluye identificación, abstracts, desarrollo, conclusión, reflexión, matrices y referencias.",
        "La propuesta cubre los indicadores seleccionados del perfil de egreso con evidencia planificada.",
        "Abstract, conclusion and reflection communicate connected, project-specific ideas in English.",
    ]
    for row_index, value in enumerate(rubric_evidence, start=1):
        replace_paragraph(doc.tables[8].cell(row_index, 3).paragraphs[0], value)

    # La plantilla combina un salto de página manual con el salto propio del
    # encabezado "Referencias", lo que genera una hoja completamente vacía.
    references_heading = next(
        paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "Referencias"
    )
    previous_element = references_heading._p.getprevious()
    if previous_element is not None:
        for page_break in previous_element.xpath(".//w:br[@w:type='page']"):
            page_break.getparent().remove(page_break)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_in_runs(paragraph, {"RUAHTONE": "FIX & GO"})
        for paragraph in section.footer.paragraphs:
            replace_in_runs(paragraph, {"PTY4614": "APT122"})

    doc.core_properties.title = "Autoevaluación Definición Proyecto APT - Fix & Go"
    doc.core_properties.subject = "APT122 - Evidencia individual 1.3"
    doc.core_properties.author = "Benjamín Olmedo"
    doc.core_properties.keywords = "APT122, Proyecto APT, Fix & Go, autoevaluación, fase 1"
    doc.core_properties.comments = "Documento individual de Benjamín Olmedo basado en la definición oficial consolidada de Fix & Go."
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1]
        paragraph._element.getparent().remove(paragraph._element)
    doc.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--template-11", type=Path, required=True)
    parser.add_argument("--template-12", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    out_11 = args.output_dir / "Olmedo_Benjamin_1.1_APT122_AutoevaluacionCompetenciasFase1.docx"
    out_12 = args.output_dir / "Olmedo_Benjamin_1.2_APT122_DiarioReflexionFase1.docx"
    out_13 = args.output_dir / "Olmedo_Benjamin_1.3_APT122_AutoevaluacionFase1.docx"
    if not out_13.exists():
        raise FileNotFoundError(f"No se encontró la evidencia 1.3 existente: {out_13}")

    fill_11(args.template_11, out_11)
    fill_12(args.template_12, out_12)
    fill_13(out_13)
    print(out_11)
    print(out_12)
    print(out_13)


if __name__ == "__main__":
    main()
